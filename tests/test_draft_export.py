"""Draft export + measurement: metrics, DOCX, export view.

Offline: tmp databases/uploads/exports. Never touches data/hakiscribe.db
or the real storage/ folders.
"""

import json
from pathlib import Path

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

import backend.core.config as config
from backend.repositories import exports_repo, flags_repo, segments_repo
from backend.repositories import corrections_repo
from backend.services import ingestion_service
from backend.services.evaluation_service import (
    cer,
    compare_ai_vs_reviewed,
    compare_to_reference,
    normalize_text,
    review_time,
    speaker_attribution_accuracy,
    traceability_completeness,
    wer,
)
from backend.services.export_service import (
    FOOTER_DISCLAIMER,
    WATERMARK_TEXT,
    generate_draft_docx,
)
from backend.services.ingestion_service import validate_and_store_upload
from backend.services.review_service import edit_segment, run_review
from backend.services.transcription_service import transcribe_recording
from database.db import get_connection
from database.schema import init_db
from frontend.views import transcribe_view

ROOT = Path(__file__).resolve().parent.parent
APP_PATH = ROOT / "frontend" / "app.py"

FORBIDDEN = ["certified by", "official record", "approved"]


def _recording_with_segments(tmp_path, texts=None):
    recording = validate_and_store_upload(
        "hearing.mp3",
        b"synthetic-audio",
        "Reviewer Amina",
        upload_dir=tmp_path / "uploads",
        db_path=tmp_path / "t.db",
    )
    db = tmp_path / "t.db"
    lines = texts or [
        "Court is now in session.",
        "We must attend the mention.",
        "We adjourn until Friday.",
    ]
    segments_repo.create_segments(
        recording["id"],
        [
            {
                "idx": idx,
                "start_sec": float(idx * 5),
                "end_sec": float(idx * 5 + 4),
                "speaker": "Speaker A" if idx % 2 == 0 else "Speaker B",
                "speaker_label": "A" if idx % 2 == 0 else "B",
                "text_original": text,
            }
            for idx, text in enumerate(lines)
        ],
        db_path=db,
    )
    return recording["id"], db


def _visible_text(app: AppTest) -> str:
    chunks = []
    for name in ("markdown", "caption", "info", "success", "warning", "text"):
        try:
            items = list(getattr(app, name))
        except Exception:
            continue
        for element in items:
            try:
                chunks.append(str(element.value))
            except Exception:
                pass
    return "\n".join(chunks)


def test_wer_cer_known_examples():
    assert wer("hello world", "hello world") == 0.0
    assert wer("", "") == 0.0
    assert cer("", "") == 0.0
    assert wer("", "hello") == 1.0
    assert cer("", "abc") == 1.0
    assert wer("the accused was present", "the accused present") == 0.25
    assert wer("a b", "a b c d") == 1.0
    assert wer("a b c", "a x c") == pytest.approx(1 / 3)
    assert wer("Hello, World!", "hello world") == 0.0
    assert cer("abc", "adc") == pytest.approx(1 / 3)
    assert cer("kitten", "sitting") == 0.5


def test_normalisation_keeps_diacritics_and_swahili():
    assert (
        normalize_text("Mheshimiwa, mlalamikaji yupo kotini leo.")
        == "mheshimiwa mlalamikaji yupo kotini leo"
    )
    assert normalize_text("naïve café") == "naïve café"
    assert wer("naïve café", "naive cafe") == 1.0


def test_ai_vs_reviewed_after_one_edit(tmp_path):
    recording_id, db = _recording_with_segments(tmp_path)
    before = compare_ai_vs_reviewed(recording_id, db_path=db)
    assert before["values"]["wer"] == 0.0
    assert before["values"]["segments_changed_pct"] == 0.0
    assert before["reason"] is None
    segment = segments_repo.list_segments(recording_id, db_path=db)[1]
    edit_segment(
        segment["id"],
        "We must attend the mention today.",
        "Reviewer Amina",
        "Heard the last word.",
        segment["text_current"],
        db_path=db,
    )
    after = compare_ai_vs_reviewed(recording_id, db_path=db)
    assert after["values"]["wer"] > 0
    assert after["values"]["cer"] > 0
    assert after["values"]["segments_changed_pct"] == 33.33
    assert after["sample_size"] == 3
    assert after["how_computed"]
    assert after["reason"] is None


def test_speaker_attribution_known_answer(tmp_path):
    recording_id, db = _recording_with_segments(
        tmp_path, texts=["First passage here.", "Second passage here."]
    )
    segments = segments_repo.list_segments(recording_id, db_path=db)
    assert len(segments) == 2
    conn = get_connection(db)
    try:
        conn.execute(
            "UPDATE segments SET start_sec = ?, end_sec = ? WHERE id = ?",
            (0.0, 8.0, segments[0]["id"]),
        )
        conn.execute(
            "UPDATE segments SET start_sec = ?, end_sec = ? WHERE id = ?",
            (8.0, 20.0, segments[1]["id"]),
        )
        conn.commit()
    finally:
        conn.close()
    reference = [
        {"start_ms": 0, "end_ms": 10000, "voice": "Voice One",
         "text": "First passage here."},
        {"start_ms": 10000, "end_ms": 20000, "voice": "Voice Two",
         "text": "Second passage here."},
    ]
    result = speaker_attribution_accuracy(
        recording_id, reference, db_path=db
    )
    assert result["values"]["attribution_pct"] == pytest.approx(90.0)
    assert result["values"]["boundary_mae_ms"] == pytest.approx(1000.0)
    assert result["sample_size"] == 2
    assert result["reason"] is None


def test_compare_to_reference_accuracy(tmp_path):
    recording_id, db = _recording_with_segments(
        tmp_path, texts=["court is now in session", "we adjourn"]
    )
    same = compare_to_reference(
        recording_id,
        [{"text": "court is now in session"}, {"text": "we adjourn"}],
        db_path=db,
    )
    assert same["values"]["wer"] == 0.0
    assert same["sample_size"] == 2
    changed = compare_to_reference(
        recording_id,
        [{"text": "court is now in session"}, {"text": "we adjourn today"}],
        db_path=db,
    )
    assert changed["values"]["wer"] == pytest.approx(1 / 8)


def test_traceability_completeness(tmp_path):
    recording_id, db = _recording_with_segments(tmp_path)
    assert traceability_completeness(recording_id, db_path=db)["reason"] == (
        "no corrections recorded"
    )
    segment = segments_repo.list_segments(recording_id, db_path=db)[0]
    corrections_repo.create_correction(
        segment["id"], None, "a", "b", "Reviewer Amina", "EDIT_TEXT",
        reason="Why here.", db_path=db,
    )
    corrections_repo.create_correction(
        segment["id"], None, "b", "c", "Reviewer Amina", "ACCEPT_SUGGESTION",
        db_path=db,
    )
    partial = traceability_completeness(recording_id, db_path=db)
    assert partial["values"]["completeness_pct"] == 50.0
    assert partial["sample_size"] == 2
    corrections_repo.create_correction(
        segment["id"], None, "c", "d", "Reviewer B", "DISMISS_FLAG",
        reason="False alarm.", db_path=db,
    )
    full = traceability_completeness(recording_id, db_path=db)
    assert full["values"]["completeness_pct"] == 66.67


def test_review_time_idle_gap_exclusion(tmp_path):
    recording_id, db = _recording_with_segments(tmp_path)
    assert review_time(recording_id, db_path=db)["reason"] == (
        "no review events recorded"
    )
    init_db(db)
    conn = get_connection(db)
    try:
        for stamp in (
            "2026-10-28 10:00:00",
            "2026-10-28 10:00:30",
            "2026-10-28 10:05:00",
            "2026-10-28 10:05:10",
        ):
            conn.execute(
                """INSERT INTO review_events (recording_id, reviewer, event, at)
                   VALUES (?, ?, ?, ?)""",
                (recording_id, "Reviewer Amina", "decision", stamp),
            )
        conn.commit()
    finally:
        conn.close()
    segments = segments_repo.list_segments(recording_id, db_path=db)
    for segment in segments[:2]:
        flag_id = flags_repo.create_flag(
            segment["id"], "UNCLEAR_PASSAGE", "LOW", "Check.", db_path=db
        )
        flags_repo.update_flag_status(
            flag_id, "RESOLVED", resolved_by="Reviewer Amina", db_path=db
        )
    result = review_time(recording_id, db_path=db)
    assert result["values"]["total_seconds"] == 40.0
    assert result["values"]["seconds_per_resolved_flag"] == 20.0
    assert result["sample_size"] == 4


def _transcribed_recording(tmp_path):
    recording = validate_and_store_upload(
        "hearing.mp3",
        b"synthetic-audio",
        "Reviewer Amina",
        upload_dir=tmp_path / "uploads",
        db_path=tmp_path / "t.db",
    )
    db = tmp_path / "t.db"
    transcribe_recording(recording["id"], force_mock=True, db_path=db)
    return recording["id"], db


def test_docx_full_content(tmp_path):
    recording_id, db = _transcribed_recording(tmp_path)
    segments = segments_repo.list_segments(recording_id, db_path=db)
    edit_segment(
        segments[0]["id"],
        "Court is now in session today.",
        "Reviewer Amina",
        "Heard the last word.",
        segments[0]["text_current"],
        db_path=db,
    )
    flags_repo.create_flag(
        segments[1]["id"],
        "NAME_CANDIDATE",
        "HIGH",
        "Check this name.",
        evidence="mlalamikaji",
        db_path=db,
    )
    result = generate_draft_docx(
        recording_id, "Reviewer Amina", db_path=db,
        export_dir=tmp_path / "exports",
    )
    assert result["docx_filename"].startswith("hakiscribe_draft_")
    assert result["docx_filename"].endswith(".docx")
    assert result["unresolved_count"] == 1
    docx_path = Path(result["docx_path"])
    assert docx_path.is_file()
    document = Document(str(docx_path))

    header_xml = document.sections[0].header._element.xml
    assert WATERMARK_TEXT in header_xml
    assert "NOT CERTIFIED" in header_xml
    footer_text = "\n".join(
        paragraph.text for paragraph in document.sections[0].footer.paragraphs
    )
    assert FOOTER_DISCLAIMER in footer_text

    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "MOCK TRANSCRIPT" in body
    assert "[00:00-00:04]" in body
    assert "[EDITED]" in body
    assert "[UNRESOLVED: Check this name.]" in body
    assert "Appendix A" in body
    assert "Appendix B" in body
    assert "Appendix C" in body
    assert "How computed" in body
    assert "Synthetic demo content only" in body
    assert "Reviewer Amina" in body
    lowered = (body + "\n" + header_xml + "\n" + footer_text).lower()
    for phrase in FORBIDDEN:
        assert phrase not in lowered, phrase

    row = exports_repo.get_latest_export_by_recording(recording_id, db_path=db)
    assert row["id"] == result["export_id"]
    assert row["wer"] > 0
    assert row["segments_changed_pct"] == pytest.approx(20.0)
    assert row["flags_total"] == 1
    assert row["flags_resolved"] == 0
    assert row["traceability_pct"] == 100.0
    assert row["review_seconds"] is None
    assert row["ref_wer"] is None


def test_export_without_transcript_refused(tmp_path):
    recording = validate_and_store_upload(
        "hearing.mp3",
        b"synthetic-audio",
        "Reviewer Amina",
        upload_dir=tmp_path / "uploads",
        db_path=tmp_path / "t.db",
    )
    db = tmp_path / "t.db"
    with pytest.raises(ValueError, match="Transcribe"):
        generate_draft_docx(
            recording["id"], "Reviewer Amina", db_path=db,
            export_dir=tmp_path / "exports",
        )
    recording_id, _ = _transcribed_recording(tmp_path / "other")
    with pytest.raises(ValueError, match="reviewer is required"):
        generate_draft_docx(
            recording_id, "  ", db_path=tmp_path / "other" / "t.db",
            export_dir=tmp_path / "exports",
        )


def test_metrics_none_without_reference_or_events(tmp_path):
    recording_id, db = _recording_with_segments(tmp_path)
    reference = compare_to_reference(recording_id, db_path=db)
    assert reference["values"] == {}
    assert "no reference" in reference["reason"]
    speaker = speaker_attribution_accuracy(recording_id, db_path=db)
    assert speaker["values"] == {}
    assert "no reference" in speaker["reason"]
    assert review_time(recording_id, db_path=db)["values"] == {}


def test_apptest_export_view(tmp_path, monkeypatch):
    db_path = tmp_path / "hakiscribe.db"
    for name in ("uploads", "exports"):
        (tmp_path / name).mkdir()
    monkeypatch.setenv("HAKISCRIBE_DB_PATH", str(db_path))
    monkeypatch.setenv("HAKISCRIBE_UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setenv("HAKISCRIBE_EXPORT_DIR", str(tmp_path / "exports"))
    monkeypatch.setattr(config, "is_configured", lambda: False)
    demo_wav = tmp_path / "demo.wav"
    demo_wav.write_bytes(b"synthetic-demo-audio")
    monkeypatch.setattr(ingestion_service, "DEMO_RECORDING_PATH", demo_wav)
    monkeypatch.setattr(transcribe_view, "is_configured", lambda: False)

    app = AppTest.from_file(str(APP_PATH))
    app.run()
    assert not app.exception
    app.text_input(key="reviewer_name").set_value("Reviewer Amina")
    app.run()
    app.button(key="load_demo_recording").click()
    app.run()
    app.button(key="transcribe_recording").click()
    app.run()
    app.run()
    assert not app.exception

    recording_id = app.session_state["recording_id"]
    run_review(recording_id, force_mock=True, db_path=db_path)

    app.session_state["current_step"] = "export"
    app.run()
    assert not app.exception
    blob = _visible_text(app)
    assert "not certified" in blob.lower()
    assert "Human changes (AI vs reviewed)" in blob
    assert "Not available: no reference script" in blob
    assert "Unresolved flags: 2" in blob
    assert "unresolved flags" in blob.lower()

    app.button(key="generate_draft_docx").click()
    app.run()
    assert not app.exception
    assert "Draft saved:" in _visible_text(app)
    assert len(list(app.download_button)) >= 1
    app.run()
    assert not app.exception
    assert len(list(app.download_button)) >= 1
    assert list((tmp_path / "exports").glob("*.docx"))
