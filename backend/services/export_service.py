"""Draft record export: watermarked DOCX plus recorded metrics.

A draft is never certified here: every page header carries a DRAFT
watermark and banner, the footer disclaims certification, and the
document states its synthetic demo origin. The phrases "certified by",
"official record" and "approved" never appear in generated documents.
"""

from datetime import datetime, timezone
from pathlib import Path

from backend.repositories import (
    corrections_repo,
    exports_repo,
    flags_repo,
    review_events_repo,
    segments_repo,
)
from backend.services import evaluation_service
from backend.services.ingestion_service import get_recording_or_raise
from backend.services.review_service import get_review_summary, get_speaker_roles

WATERMARK_TEXT = "DRAFT - NOT CERTIFIED"
HEADER_BANNER = "DRAFT \u2014 NOT A CERTIFIED RECORD"
FOOTER_DISCLAIMER = (
    "Draft prepared with AI assistance for review by authorised court staff. "
    "Not a certified record."
)
SYNTHETIC_STATEMENT = (
    "Synthetic demo content only. This draft was produced from a synthetic "
    "demonstration recording with no real litigants. It is not a certified "
    "record."
)
EXPORT_MIME = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document"
)


def _add_draft_watermark(header) -> None:
    """Insert a diagonal VML text-path watermark into the page header."""
    from docx.oxml import parse_xml

    paragraph = header.add_paragraph()
    run = paragraph.add_run()
    pict = parse_xml(
        '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:w10="urn:schemas-microsoft-com:office:word">'
        '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" '
        'adj="10800" path="m@7,l@8,m@5,21600l@6,21600e">'
        '<v:path textpathok="t" o:connecttype="custom" '
        'o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" '
        'o:connectangles="270,180,90,0"/>'
        '<v:textpath on="t" fitshape="t"/>'
        '<o:lock v:ext="edit" text="t" shapetype="t"/>'
        "</v:shapetype>"
        '<v:shape id="hakiscribeDraftWatermark" type="#_x0000_t136" '
        'style="position:absolute;margin-left:0;margin-top:0;'
        'width:468pt;height:96pt;rotation:315;'
        'mso-position-horizontal:center;mso-position-vertical:center;'
        'mso-position-vertical-relative:page" '
        'o:spid="_x0000_s1026" fillcolor="#d9d9d9" stroked="f">'
        '<v:fill opacity=".35"/>'
        '<v:textpath style="font-family:&quot;Calibri&quot;" '
        f'string="{WATERMARK_TEXT}"/>'
        '<w10:wrap anchorx="page" anchory="page"/>'
        "</v:shape>"
        "</w:pict>"
    )
    run._r.append(pict)


def _format_clock(seconds) -> str:
    """Render seconds as mm:ss (00:00 when unknown)."""
    total = int(seconds or 0)
    return f"{total // 60:02d}:{total % 60:02d}"


def _metric_line(name: str, envelope: dict) -> str:
    """Render one metric envelope as a single document line."""
    if not envelope["values"]:
        return f"{name}: Not available: {envelope['reason']}"
    parts = ", ".join(
        f"{key}={value}" for key, value in envelope["values"].items()
    )
    return (
        f"{name}: {parts} (n={envelope['sample_size']}). "
        f"How computed: {envelope['how_computed']}"
    )


def _voice_mode_for(recording_id: int, db_path) -> str:
    """Latest voice-attempt mode for a recording, else the configured mode."""
    from backend.repositories import api_calls_repo

    for call in api_calls_repo.list_recent_api_calls(
        limit=50, db_path=db_path
    ):
        if (
            call["recording_id"] == recording_id
            and call["endpoint_path"] == "/asr/transcribe"
        ):
            return call["mode"]
    from backend.core.config import get_voice_mode

    return get_voice_mode()


def generate_draft_docx(
    recording_id: int,
    reviewer: str,
    db_path: str | Path | None = None,
    export_dir: str | Path = "storage/exports",
) -> dict:
    """Build the watermarked draft DOCX, save it, and record the export.

    Raises ValueError when the reviewer is empty or the recording has no
    transcript. Unresolved flags are allowed; they are listed inline and
    in Appendix B. Returns filenames, the export id, and metric envelopes.
    """
    if not reviewer or not reviewer.strip():
        raise ValueError("reviewer is required")
    reviewer = reviewer.strip()
    recording = get_recording_or_raise(recording_id, db_path=db_path)
    segments = segments_repo.list_segments(recording_id, db_path=db_path)
    if not segments:
        raise ValueError("Transcribe the recording before export.")
    flags = flags_repo.list_flags_by_recording(recording_id, db_path=db_path)
    corrections = corrections_repo.list_corrections_by_recording(
        recording_id, db_path=db_path
    )
    roles = get_speaker_roles(recording_id, db_path=db_path)
    summary = get_review_summary(recording_id, db_path=db_path)
    events = review_events_repo.list_events_by_recording(
        recording_id, db_path=db_path
    )
    ai_metrics = evaluation_service.compare_ai_vs_reviewed(
        recording_id, db_path=db_path
    )
    ref_metrics = evaluation_service.compare_to_reference(
        recording_id, None, db_path=db_path
    )
    speaker_metrics = evaluation_service.speaker_attribution_accuracy(
        recording_id, None, db_path=db_path
    )
    trace_metrics = evaluation_service.traceability_completeness(
        recording_id, db_path=db_path
    )
    time_metrics = evaluation_service.review_time(recording_id, db_path=db_path)
    unresolved = [flag for flag in flags if flag["status"] in ("OPEN", "RETAINED")]
    resolved = sum(1 for flag in flags if flag["status"] == "RESOLVED")

    source = (recording.get("transcription_source") or "mock").lower()
    is_mock = source != "real"
    from docx import Document

    document = Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = HEADER_BANNER
    _add_draft_watermark(section.header)
    section.footer.paragraphs[0].text = FOOTER_DISCLAIMER

    document.add_heading(
        "Draft Hearing Transcript (DRAFT - NOT CERTIFIED)", level=0
    )
    if is_mock:
        mock_note = document.add_paragraph()
        mock_note.add_run("MOCK TRANSCRIPT \u2014 SYNTHETIC DEMO ONLY").bold = True

    duration = recording.get("duration_sec")
    if duration is None:
        span = max((segment["end_sec"] or 0) for segment in segments)
        duration_text = f"transcript spans 0\u2013{span:.1f}s"
    else:
        duration_text = f"{duration:.1f}s"
    reviewers = sorted(
        {row["reviewer_name"] for row in corrections if row.get("reviewer_name")}
        | {event["reviewer"] for event in events if event.get("reviewer")}
    )
    if summary is None:
        review_line = "AI review not run"
    else:
        review_line = (
            f"{summary['model']} (prompt {summary['prompt_version']}, "
            f"mode {summary['mode']})"
        )
    for line in (
        f"Recording: {recording['filename']}",
        f"SHA-256: {recording['sha256']}",
        f"Duration: {duration_text}",
        "Transcript source: "
        + ("MOCK (synthetic demo)" if is_mock else "REAL Meta transcription"),
        f"Voice model: {recording.get('transcription_model') or 'unknown'} "
        f"(mode {_voice_mode_for(recording_id, db_path)})",
        f"Review model: {review_line}",
        f"Reviewers: {', '.join(reviewers) if reviewers else 'None yet'}",
        f"Generated: {datetime.now(timezone.utc):%Y-%m-%d %H:%M UTC} "
        f"by {reviewer}",
        f"Corrections: {len(corrections)}",
        f"Unresolved flags: {len(unresolved)}",
    ):
        document.add_paragraph(line)

    document.add_heading("Transcript", level=1)
    for segment in segments:
        label = segment.get("speaker_label")
        speaker_display = (
            roles.get(label) if label and roles.get(label) else segment["speaker"]
        )
        line = (
            f"[{_format_clock(segment['start_sec'])}-"
            f"{_format_clock(segment['end_sec'])}] "
            f"{speaker_display}: {segment['text_current']}"
        )
        if (segment["text_current"] or "") != (segment["text_original"] or ""):
            line += " [EDITED]"
        for flag in unresolved:
            if flag["segment_id"] == segment["id"]:
                line += f" [UNRESOLVED: {flag['message']}]"
        document.add_paragraph(line)

    document.add_heading("Appendix A \u2014 Correction log", level=1)
    if not corrections:
        document.add_paragraph("No corrections recorded.")
    by_segment = {segment["id"]: segment for segment in segments}
    for correction in corrections:
        segment = by_segment.get(correction["segment_id"], {})
        clock = (
            f"[{_format_clock(segment.get('start_sec'))}-"
            f"{_format_clock(segment.get('end_sec'))}]"
        )
        document.add_paragraph(
            f"Segment {segment.get('idx', '?')} {clock}: "
            f"{correction['reviewer_name']} at {correction['created_at']}: "
            f"{correction['decision']} [{correction['field']}]: "
            f"{correction['before_text']} \u2192 {correction['after_text']}. "
            f"Reason: {correction['reason'] or '\u2014'}"
        )

    document.add_heading("Appendix B \u2014 Unresolved issues", level=1)
    if not unresolved:
        document.add_paragraph("None \u2014 every flag has a decision.")
    for flag in unresolved:
        clock = (
            f"[{_format_clock(flag.get('start_sec'))}-"
            f"{_format_clock(flag.get('end_sec'))}]"
        )
        document.add_paragraph(
            f"Segment {flag.get('segment_idx', '?')} {clock} "
            f"{flag['flag_type']} ({flag['severity']}): {flag['message']} "
            f"Quote: {flag['evidence']}"
        )

    document.add_heading("Appendix C \u2014 Evaluation metrics", level=1)
    for name, envelope in (
        ("Human changes (AI vs reviewed)", ai_metrics),
        ("Accuracy vs reference", ref_metrics),
        ("Speaker attribution", speaker_metrics),
        ("Traceability", trace_metrics),
        ("Review time", time_metrics),
    ):
        document.add_paragraph(_metric_line(name, envelope))
    document.add_paragraph(SYNTHETIC_STATEMENT)

    out_dir = Path(export_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    filename = f"hakiscribe_draft_{recording_id}_{stamp}.docx"
    docx_path = out_dir / filename
    document.save(str(docx_path))

    ai_values = ai_metrics["values"]
    ref_values = ref_metrics["values"]
    speaker_values = speaker_metrics["values"]
    trace_values = trace_metrics["values"]
    time_values = time_metrics["values"]
    export_id = exports_repo.create_export(
        recording_id,
        filename,
        wer=ai_values.get("wer"),
        cer=ai_values.get("cer"),
        segments_changed_pct=ai_values.get("segments_changed_pct"),
        flags_total=len(flags),
        flags_resolved=resolved,
        generated_by=reviewer,
        ref_wer=ref_values.get("wer"),
        ref_cer=ref_values.get("cer"),
        speaker_attribution_pct=speaker_values.get("attribution_pct"),
        boundary_error_ms=speaker_values.get("boundary_mae_ms"),
        traceability_pct=trace_values.get("completeness_pct"),
        review_seconds=time_values.get("total_seconds"),
        db_path=db_path,
    )
    return {
        "recording_id": recording_id,
        "docx_filename": filename,
        "docx_path": str(docx_path),
        "export_id": export_id,
        "unresolved_count": len(unresolved),
        "metrics": {
            "ai_vs_reviewed": ai_metrics,
            "reference": ref_metrics,
            "speaker": speaker_metrics,
            "traceability": trace_metrics,
            "review_time": time_metrics,
        },
    }
